import PySimpleGUI as sg
import openpyxl
import docx
import socket


def tela():
    global tel_dir, msg_dir
    # 'Layout'
    layout = [
        [sg.Text('Mensagem'), sg.In(key='-MSG-'), sg.FileBrowse(file_types=(("Text Files", "*.docx"),))],
        [sg.Text('Telefone    '), sg.In(key='-TEL-'), sg.FileBrowse(file_types=(("Text Files", "*.xlsx"),))],
        [sg.Button('Enviar'), sg.Button('Sair')]
    ]
    # 'Janela'
    window = sg.Window('WppSender - CreatedBy Vini Almeida @v.alma_br', layout)

    while True:
        event, values = window.read()
        if event == 'Sair' or event is None:
            tel_dir = ''
            msg_dir = ''
            exit()
        elif event == 'Enviar':
            tel_dir = values['-TEL-']
            msg_dir = values['-MSG-']
        window.close()
        return tel_dir, msg_dir


def get_text(filename):
    try:
        doc = docx.Document(filename)
        fulltext = []
        for para in doc.paragraphs:
            fulltext.append(para.text)
        return fulltext  # '\n'.join(fullText)
    except Exception as erro:
        errpop2 = [
            [sg.Text('Mensagem .docx não encontrada ou inválida.')],
            [sg.Text('Consulte a documentação')],
            [sg.Button('Fechar')]
        ]
        popup2 = sg.Window('WppSender - Arquivo inválido', errpop2)
        while True:
            event, values = popup2.read()
            if event == 'Fechar' or event is None:
                exit()


def get_excel(filename):
    try:
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        #ultlin = ws.max_row #UltLin = len(sheet['A'])
        vector = []
        for row in ws.values:
            vector.append(row)

        del vector[0]
        return vector
    except Exception as erro:
        errpop = [
            [sg.Text('Telefones .xlsx não encontrado ou inválido: {}')],
            [sg.Text('Consulte a documentação')],
            [sg.Button('Fechar')]
        ]
        popup = sg.Window('WppSender - Arquivo inválido', errpop)
        while True:
            event, values = popup.read()
            if event == 'Fechar' or event is None:
                exit()


def is_connected():
    try:
        # connect to the host
        socket.create_connection(("www.google.com", 80))
        return True
    except Exception:
        is_connected()


