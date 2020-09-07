from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.keys import Keys
from selenium import webdriver
from time import sleep
import PySimpleGUI as sg
import openpyxl
import os 
import docx
import socket
import logging

#logging.basicConfig(filename='LastLogWppSender.txt', level=logging.DEBUG,
#                    format='%(asctime)s - %(levelname)s - %(message)s')


def tela():
    global tel_dir, msg_dir
    # 'Layout'
    layout = [
        [sg.Text('Mensagem'), sg.In(key='-MSG-'), sg.FileBrowse(file_types=(("Text Files", "*.docx"),))],
        [sg.Text('Telefone    '), sg.In(key='-TEL-'), sg.FileBrowse(file_types=(("Text Files", "*.xlsx"),))],
        [sg.Button('Enviar'), sg.Button('Sair')]
    ]
    # 'Janela'
    window = sg.Window('WppSender - CreatedBy Vini Almeida @v.alma_br', layout)

    while True:
        event, values = window.read()
        if event == 'Sair' or event is None:
            tel_dir = ''
            msg_dir = ''
            exit()
        elif event == 'Enviar':
            tel_dir = values['-TEL-']
            msg_dir = values['-MSG-']
        window.close()
        return tel_dir, msg_dir


def get_text(filename):
    try:
        doc = docx.Document(filename)
        fulltext = []
        for para in doc.paragraphs:
            fulltext.append(para.text)
        return fulltext  # '\n'.join(fullText)
    except Exception as errpop2:
        errpop2 = [
            [sg.Text('Mensagem .docx não encontrada ou inválida.')],
            [sg.Text('Consulte a documentação')],
            [sg.Button('Fechar')]
        ]
        popup2 = sg.Window('WppSender - Arquivo inválido', errpop2)
        while True:
            event, values = popup2.read()
            if event == 'Fechar' or event is None:
                exit()


def get_excel(filename):
    try:
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        #ultlin = ws.max_row #UltLin = len(sheet['A'])
        vector = []
        for row in ws.values:
            vector.append(row)

        del vector[0]
        return vector
    except Exception as errpop:
        errpop = [
            [sg.Text('Telefones .xlsx não encontrado ou inválido: {}')],
            [sg.Text('Consulte a documentação')],
            [sg.Button('Fechar')]
        ]
        popup = sg.Window('WppSender - Arquivo inválido', errpop)
        while True:
            event, values = popup.read()
            if event == 'Fechar' or event is None:
                exit()


def is_connected():
    try:
        # connect to the host
        socket.create_connection(("www.google.com", 80))
        return True
    except Exception:
        is_connected()


def enviar_mensagem(driver, telephone, text):
    driver.get('https://web.whatsapp.com/send?phone={}'.format(telephone))
    sleep(6)
    try:
        txt_box = driver.find_element_by_xpath('//*[@id="main"]/footer/div[1]/div[2]/div/div[2]')
        for i in range(0, len(text)):
            txt_box.send_keys(text[i])
            ActionChains(driver).key_down(Keys.SHIFT).key_down(Keys.ENTER).key_up(Keys.SHIFT).key_up(
                Keys.ENTER).perform()
        sleep(3)
        btn_enviar = driver.find_element_by_xpath('//*[@id="main"]/footer/div[1]/div[3]/button')
        btn_enviar.click()
        sleep(1)

    except Exception:
        print('Erro ao enviar para o phone no: ' + str(telephone))


def main():
    (tel_dir, msg_dir) = tela()  # ABRE TELA
    telephones = []
    names = []
    text = get_text(msg_dir)  # 'LER MENSAGEM DOCX'
    vector = get_excel(tel_dir)  # 'LER TELEFONES XLSX'

    path = os.getcwd().replace("\\","/")+'/chromedriver.exe'
    driver = webdriver.Chrome(executable_path= path)
    driver.get("http://web.whatsapp.com")
    sleep(10)

    for i in range(0, len(vector)):
        telephones.append(vector[i][1])
        names.append(vector[i][0])

        try:
            header = "Olá, {}! Tudo bem?".format(names[i])
            del text[0]
            text.insert(0, header)
            enviar_mensagem(driver, telephones[i], text)

        except Exception:
            sleep(10)
            is_connected()


if __name__ == '__main__':
    main()
